"""
Convert non-PDF formats (images, Word, Excel) into PDF bytes so the rest of
the pipeline (which expects PDFs) can handle them unchanged.
"""

import io
from pathlib import Path

import fitz  # PyMuPDF – already installed

SUPPORTED_EXTENSIONS = {
    ".pdf": "PDF",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".tiff": "image",
    ".tif": "image",
    ".bmp": "image",
    ".docx": "word",
    ".doc": "word",
    ".xlsx": "excel",
    ".xls": "excel",
}


def is_supported(filename: str) -> bool:
    ext = Path(filename).suffix.lower()
    return ext in SUPPORTED_EXTENSIONS


def get_format(filename: str) -> str:
    return SUPPORTED_EXTENSIONS.get(Path(filename).suffix.lower(), "unknown")


def to_pdf_bytes(content: bytes, filename: str) -> bytes:
    """Convert file bytes to PDF bytes. Returns original if already PDF."""
    fmt = get_format(filename)
    if fmt == "PDF":
        return content
    if fmt == "image":
        return _image_to_pdf(content)
    if fmt == "word":
        return _word_to_pdf(content)
    if fmt == "excel":
        return _excel_to_pdf(content)
    raise ValueError(f"Unsupported format: {filename}")


def _image_to_pdf(content: bytes) -> bytes:
    """Wrap a single image as a PDF page using PyMuPDF."""
    img_doc = fitz.open(stream=content, filetype="*")
    pdf_bytes = img_doc.convert_to_pdf()
    img_doc.close()
    return pdf_bytes


def _word_to_pdf(content: bytes) -> bytes:
    """Extract text from a .docx file and build a simple text PDF."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    except Exception:
        paragraphs = ["[Word document — could not extract text]"]

    return _text_to_pdf("\n\n".join(paragraphs))


def _excel_to_pdf(content: bytes) -> bytes:
    """Extract each sheet from an Excel file and build a text PDF."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        pages_text = []
        for sheet in wb.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cells = "\t".join(str(v) if v is not None else "" for v in row)
                if cells.strip():
                    rows.append(cells)
            if rows:
                pages_text.append(f"Sheet: {sheet.title}\n\n" + "\n".join(rows))
        text = "\n\n---\n\n".join(pages_text) if pages_text else "[Empty workbook]"
    except Exception:
        text = "[Excel document — could not extract data]"

    return _text_to_pdf(text)


def _text_to_pdf(text: str) -> bytes:
    """Render plain text into a PDF using PyMuPDF."""
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)  # A4 points
    rect = fitz.Rect(50, 50, 545, 792)
    page.insert_textbox(rect, text, fontsize=10, fontname="helv", color=(0, 0, 0))
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes
